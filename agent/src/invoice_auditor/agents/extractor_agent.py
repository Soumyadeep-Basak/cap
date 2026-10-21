
"""LangGraph workflow for extracting text from invoice documents."""
from decimal import Decimal
import logging
import uuid
from pathlib import Path
from typing import Any, TypedDict

import pdfplumber
import pytesseract
from docx import Document
from pdf2image import convert_from_path
from langgraph.graph import END, START, StateGraph

from invoice_auditor.db.session import SessionLocal
from invoice_auditor.services.invoice_audit_service import (
    mark_processing_failed,
    save_extraction_result,
)

logger = logging.getLogger(__name__)

OCR_LANGUAGES = "eng+spa+deu"


class ExtractorState(TypedDict, total=False):
    invoice: dict[str, Any]
    raw_extracted_text: str
    extraction_method: str
    status: str
    error: str



import re
from datetime import datetime
from decimal import InvalidOperation


def parse_amount(value: str) -> Decimal | None:
    """Convert a monetary value containing optional currency symbols to Decimal."""
    cleaned = re.sub(
        r"[^\d.\-]",
        "",
        value.strip().replace(",", ""),
    )

    if not cleaned:
        return None

    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def parse_invoice_text(raw_text: str) -> dict:
    """Parse common invoice headers, totals, and line items."""

    invoice_number_match = re.search(
        r"^\s*(?:"
        r"Invoice\s+(?:No\.?|Number)|"
        r"Involce|"
        r"Factura\s+(?:No\.?|Número|Numero)"
        r")\s*:?\s*(?:No\.?\s*:?\s*)?(\S+)",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    # Debit notes may expose their reference as a payment/local invoice number.
    if not invoice_number_match:
        invoice_number_match = re.search(
            r"Payment\s+Ref/Local\s+inv\s+no\s*:\s*(\S+)",
            raw_text,
            re.IGNORECASE,
        )

    invoice_date_match = re.search(
        r"^\s*(?:Invoice\s+Date|Date|Fecha)\s*:\s*"
        r"(.+?)(?=\s+Payment\s+Ref/Local\s+inv\s+no\s*:|$)",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    po_match = re.search(
        r"^\s*(?:"
        r"PO(?:\s+(?:No\.?|Reference|Number))?|"
        r"Purchase\s+Order(?:\s+(?:No\.?|Number|Reference))?|"
        r"Referencia\s+de\s+pedido"
        r")\s*:\s*(\S+)",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    total_match = re.search(
        r"^\s*(?:"
        r"Total(?:\s+Amount)?(?:\s+Due)?|"
        r"Total\s+Including\s+Tax"
        r")\s*:?\s*"
        r"([$€£₹])?\s*([\d,]+(?:\.\d{1,2})?)\s*([$€£₹])?\s*$",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )


    subtotal_match = re.search(
        r"^\s*Subtotal\s*:?\s*"
        r"([$€£₹])?\s*([\d,]+(?:\.\d{1,2})?)\s*([$€£₹])?\s*$",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    tax_match = re.search(
        r"^\s*(?:Tax|VAT|IVA)\s*(?:\((\d+(?:\.\d+)?)%\))?"
        r"\s*:?\s*([$€£₹])?\s*"
        r"([\d,]+(?:\.\d{1,2})?)\s*([$€£₹])?\s*$",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    subtotal_amount = (
        parse_amount(subtotal_match.group(2))
        if subtotal_match
        else None
    )
    tax_amount = (
        parse_amount(tax_match.group(3))
        if tax_match
        else None
    )
    tax_rate = (
        Decimal(tax_match.group(1))
        if tax_match and tax_match.group(1)
        else None
    )

    invoice_date = None
    if invoice_date_match:
        date_text = invoice_date_match.group(1).strip()

        for date_format in (
            "%d %B %Y",
            "%d %b %Y",
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%d.%m.%Y",
            "%B %d, %Y",
            "%b %d, %Y",
            "%d-%b-%Y",
        ):
            try:
                invoice_date = datetime.strptime(
                    date_text, date_format
                ).date()
                break
            except ValueError:
                continue

    currency_symbols = {
        "$": "USD",
        "€": "EUR",
        "£": "GBP",
        "₹": "INR",
    }

    currency_match = re.search(
        r"^\s*(?:Subtotal|Tax|VAT|IVA|Total(?:\s+Amount)?"
        r"(?:\s+Due)?|Total\s+Including\s+Tax)\s*:?\s*"
        r"(?:([$€£₹])|[\d,]+(?:\.\d{1,2})?\s*([$€£₹]))",
        raw_text,
        re.IGNORECASE | re.MULTILINE,
    )

    currency = None
    if currency_match:
        symbol = currency_match.group(1) or currency_match.group(2)
        currency = currency_symbols.get(symbol)

    # Recognize explicit ISO currency codes in the extracted text.
    # If multiple currencies appear, leave the invoice currency unresolved.
    currency_codes = {
        code.upper()
        for code in re.findall(
            r"\b(USD|EUR|GBP|INR|VND)\b",
            raw_text,
            re.IGNORECASE,
        )
    }

    if len(currency_codes) == 1:
        currency = next(iter(currency_codes))
    elif len(currency_codes) > 1:
        currency = None

    # Recognize both SKU-coded rows and ordinary OCR rows.
    line_items = []
    item_patterns = (
        re.compile(
            r"^\s*(SKU-\S+)\s*\|\s*(.*?)\s*\|\s*"
            r"(\d+(?:\.\d+)?)\s*\|\s*"
            r"[$€£]?\s*([\d,]+(?:\.\d{1,2})?)\s*[$€£]?\s*\|\s*"
            r"[$€£]?\s*([\d,]+(?:\.\d{1,2})?)\s*[$€£]?\s*$",
            re.IGNORECASE,
        ),
        re.compile(
            r"^\s*([A-Za-z][A-Za-z ]*?)\s+"
            r"(\d+(?:\.\d+)?)\s+"
            r"([\d,]+(?:\.\d{1,2})?)\s+"
            r"([\d,]+(?:\.\d{1,2})?)\s*$"
        ),
    )

    for line in raw_text.splitlines():
        match = item_patterns[0].match(line)
        if match:
            item_code, description, quantity, unit_price, line_total = (
                match.groups()
            )
        else:
            match = item_patterns[1].match(line)
            if not match:
                continue

            description, quantity, unit_price, line_total = match.groups()
            item_code = None

        line_items.append({
            "item_code": item_code,
            "description": description.strip(),
            "quantity": Decimal(quantity),
            "unit_price": parse_amount(unit_price),
            "line_total": parse_amount(line_total),
        })

    # Fallback for invoices that describe each line item using labelled fields:
    #
    # Description: Container Seals
    # Quantity: 200 units
    # Unit Price: $1.25
    #
    # Do not interpret the invoice subtotal as a line total.
    if not line_items:
        labelled_items = []
        current_item: dict[str, Any] = {}

        def finish_labelled_item() -> None:
            """Save a complete labelled line item, if one was extracted."""
            if not current_item.get("description"):
                return

            quantity = current_item.get("quantity")
            unit_price = current_item.get("unit_price")

            if quantity is None or unit_price is None:
                return

            explicit_line_total = current_item.get("line_total")
            line_total = (
                explicit_line_total
                if explicit_line_total is not None
                else quantity * unit_price
            )

            labelled_items.append(
                {
                    "item_code": current_item.get("item_code"),
                    "description": current_item["description"],
                    "quantity": quantity,
                    "unit_price": unit_price,
                    "line_total": line_total,
                }
            )

        for line in raw_text.splitlines():
            stripped = line.strip()

            description_match = re.match(
                r"^(?:Item\s+)?Description\s*:\s*(.+)$",
                stripped,
                re.IGNORECASE,
            )

            if description_match:
                finish_labelled_item()
                current_item = {
                    "item_code": None,
                    "description": description_match.group(1).strip(),
                }
                continue

            if not current_item:
                continue

            quantity_match = re.match(
                r"^Quantity\s*:\s*([\d,]+(?:\.\d+)?)\b",
                stripped,
                re.IGNORECASE,
            )

            if quantity_match:
                current_item["quantity"] = parse_amount(
                    quantity_match.group(1)
                )
                continue

            price_match = re.match(
                r"^(?:Unit\s+Price|Price\s+per\s+Unit)\s*:\s*"
                r"([$€£₹]?\s*[\d,]+(?:\.\d{1,2})?)",
                stripped,
                re.IGNORECASE,
            )

            if price_match:
                current_item["unit_price"] = parse_amount(
                    price_match.group(1)
                )
                continue

            line_total_match = re.match(
                r"^(?:Line\s+Total|Amount)\s*:\s*"
                r"([$€£₹]?\s*[\d,]+(?:\.\d{1,2})?)",
                stripped,
                re.IGNORECASE,
            )

            if line_total_match:
                current_item["line_total"] = parse_amount(
                    line_total_match.group(1)
                )

        finish_labelled_item()
        line_items.extend(labelled_items)
        
    vendor_name = None
    for line in raw_text.splitlines():
        candidate = line.strip()

        if (
            not candidate
            or candidate.startswith("---")
            or candidate.startswith("-")
            or "|" in candidate
            or re.match(
                r"^(Invoice|Involce|Date|Fecha|Factura|PO\b|"
                r"Purchase Order|Referencia de pedido|"
                r"Subtotal|Tax|VAT|Total)\b",
                candidate,
                re.IGNORECASE,
            )
        ):
            continue

        vendor_name = candidate
        break

    return {
        "invoice_number": (
            invoice_number_match.group(1)
            if invoice_number_match else None
        ),
        "invoice_date": invoice_date,
        "vendor_name": vendor_name,
        "po_number": po_match.group(1) if po_match else None,
        "currency": currency,
        "total_amount": (
            parse_amount(total_match.group(2))
            if total_match else None
        ),
        "line_items": line_items,
        "subtotal_amount": subtotal_amount,
        "tax_amount": tax_amount,
        "tax_rate": tax_rate,
    }


def extract_pdf(file_path: Path) -> tuple[str, str]:
    """Extract PDF text and OCR pages without usable embedded text."""
    extracted_pages = []
    methods = set()

    with pdfplumber.open(file_path) as pdf:
        page_count = len(pdf.pages)

        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""

            if len(text.strip()) >= 20:
                extracted_pages.append(text)
                methods.add("pdf_text")
                continue

            # OCR only this page when embedded text is unusable.
            images = convert_from_path(
                str(file_path),
                first_page=page_number,
                last_page=page_number,
            )

            if images:
                text = pytesseract.image_to_string(
                    images[0],
                    lang=OCR_LANGUAGES,
                )

            extracted_pages.append(text)
            methods.add("pdf_ocr")

    if not extracted_pages:
        raise ValueError("PDF contains no pages")

    raw_text = "\n\n".join(
        f"--- Page {number} ---\n{text}"
        for number, text in enumerate(extracted_pages, start=1)
    )

    if not raw_text.strip():
        raise ValueError("No text could be extracted from the PDF")

    if methods == {"pdf_text"}:
        method = "pdf_text"
    elif methods == {"pdf_ocr"}:
        method = "pdf_ocr"
    else:
        method = "pdf_text+pdf_ocr"

    return raw_text, method


def extract_docx(file_path: Path) -> tuple[str, str]:
    """Extract text from DOCX paragraphs and tables."""
    document = Document(file_path)
    sections = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            sections.append(paragraph.text)

    for table_number, table in enumerate(document.tables, start=1):
        sections.append(f"--- Table {table_number} ---")

        for row in table.rows:
            sections.append(" | ".join(cell.text for cell in row.cells))

    raw_text = "\n".join(sections)

    if not raw_text.strip():
        raise ValueError("No text could be extracted from the DOCX")

    return raw_text, "docx_text"


def extract_png(file_path: Path) -> tuple[str, str]:
    """Extract text from a PNG image using Tesseract OCR."""
    from PIL import Image

    with Image.open(file_path) as image:
        raw_text = pytesseract.image_to_string(
            image,
            lang=OCR_LANGUAGES,
        )

    if not raw_text.strip():
        raise ValueError("OCR did not extract any text from the PNG")

    return raw_text, "png_ocr"


def extract_document(file_path: Path) -> tuple[str, str]:
    """Choose the appropriate extraction method based on file extension."""
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".png":
        return extract_png(file_path)

    raise ValueError(f"Unsupported file type: {extension or '(no extension)'}")


def extract_document_node(
    state: ExtractorState,
) -> dict[str, Any]:
    """Extract the document text."""
    invoice = state["invoice"]
    file_path = Path(invoice["file_path"])

    raw_text, extraction_method = extract_document(file_path)

    logger.info(
        "Extracted text from %s using %s",
        file_path.name,
        extraction_method,
    )

    return {
        "raw_extracted_text": raw_text,
        "extraction_method": extraction_method,
        "status": "extracted",
    }


def save_extraction_node(
    state: ExtractorState,
) -> dict[str, Any]:
    """Persist the extracted text and extraction method."""
    invoice = state["invoice"]
    invoice_id = uuid.UUID(invoice["invoice_id"])

    with SessionLocal() as db:
        save_extraction_result(
            db=db,
            invoice_id=invoice_id,
            raw_extracted_text=state["raw_extracted_text"],
            extraction_method=state["extraction_method"],
            invoice={},
            extracted_raw={},
        )

    return {"status": "completed"}


def handle_extraction_failure_node(
    state: ExtractorState,
) -> dict[str, Any]:
    """Persist an extraction failure in the audit record."""
    invoice = state["invoice"]
    invoice_id = uuid.UUID(invoice["invoice_id"])
    error_message = state.get("error", "Unknown extraction error")

    with SessionLocal() as db:
        mark_processing_failed(
            db=db,
            invoice_id=invoice_id,
            error_message=error_message,
        )

    logger.error(
        "Extraction failed for %s: %s",
        invoice.get("file_path"),
        error_message,
    )

    return {"status": "failed"}


def route_extraction_result(
    state: ExtractorState,
) -> str:
    """Route extraction to persistence or failure handling."""
    return "save_extraction" if state.get("status") == "extracted" else "handle_failure"


def build_invoice_extractor_graph():
    workflow = StateGraph(ExtractorState)

    workflow.add_node("extract_document", extract_document_node)
    workflow.add_node("save_extraction", save_extraction_node)
    workflow.add_node("handle_failure", handle_extraction_failure_node)

    workflow.add_edge(START, "extract_document")

    workflow.add_conditional_edges(
        "extract_document",
        route_extraction_result,
        {
            "save_extraction": "save_extraction",
            "handle_failure": "handle_failure",
        },
    )

    workflow.add_edge("save_extraction", END)
    workflow.add_edge("handle_failure", END)

    return workflow.compile()


invoice_extractor_graph = build_invoice_extractor_graph()
